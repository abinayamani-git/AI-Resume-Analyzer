"""
Resume text extraction from PDF and DOCX files.
"""

from pathlib import Path

from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document

from app.utils.text_utils import sanitize_text


def extract_text_from_pdf(path: Path) -> str:
    """Extract plain text from a PDF using pypdf."""
    try:
        reader = PdfReader(str(path))
        chunks: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            chunks.append(page_text)
        text = "\n".join(chunks)
        return sanitize_text(text)
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not extract text from the PDF. The file may be scanned/image-only or corrupted.",
        ) from exc


def extract_text_from_docx(path: Path) -> str:
    """Extract plain text from a DOCX using python-docx."""
    try:
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # Also pull simple table text
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
        return sanitize_text(text)
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not extract text from the DOCX file. Please try another file.",
        ) from exc


def extract_resume_text(path: Path) -> str:
    """Route to the correct extractor based on file extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(path)
    elif ext == ".docx":
        text = extract_text_from_docx(path)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Use PDF or DOCX.")

    if not text or len(text.strip()) < 40:
        raise HTTPException(
            status_code=422,
            detail="Resume text is too short or empty after extraction. Please upload a text-based resume.",
        )
    return text
