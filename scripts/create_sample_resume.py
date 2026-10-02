"""
Create a sample DOCX resume for local testing.
Run from backend folder after installing dependencies:
  python ../scripts/create_sample_resume.py
"""

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "uploads" / "Sample_Resume_Demo.docx"


def main():
    doc = Document()
    doc.add_heading("Abinaya Demo", level=1)
    doc.add_paragraph("abinaya.demo@email.com | +91 98765 43210 | Chennai, India")
    doc.add_paragraph("linkedin.com/in/abinaya-demo | github.com/abinaya-demo")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Motivated software developer with hands-on experience in Python, FastAPI, React, "
        "and Machine Learning. Built practical projects involving resume analysis, REST APIs, "
        "and modern web dashboards. Seeking opportunities in AI/ML and software development."
    )

    doc.add_heading("Skills", level=2)
    doc.add_paragraph(
        "Python, JavaScript, React, FastAPI, SQL, MongoDB, Git, GitHub, Machine Learning, "
        "NLP, Docker, REST APIs, Pytest"
    )

    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Software Development Intern | Demo Tech Solutions")
    doc.add_paragraph("Jan 2025 – Jun 2025")
    doc.add_paragraph(
        "• Developed REST APIs with FastAPI and integrated React dashboards for internal tools."
    )
    doc.add_paragraph(
        "• Automated API testing workflows and improved response validation coverage by 30%."
    )
    doc.add_paragraph(
        "• Collaborated with the team using Git/GitHub and documented technical architecture."
    )

    doc.add_heading("Projects", level=2)
    doc.add_paragraph("AI Resume Analyzer")
    doc.add_paragraph(
        "Built a full-stack resume analysis application using Python, FastAPI, and React. "
        "Implemented PDF/DOCX parsing, skill detection, estimated ATS scoring, and job-match insights."
    )
    doc.add_paragraph("Technologies: Python, FastAPI, React, Machine Learning, NLP")

    doc.add_paragraph("Job Match Insight Tool")
    doc.add_paragraph(
        "Created a keyword overlap analyzer that compares resume skills with job descriptions "
        "and highlights missing skills with recommended improvements."
    )
    doc.add_paragraph("Technologies: Python, SQL, Git")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Tech in Computer Science")
    doc.add_paragraph("Demo Institute of Technology | 2026")

    doc.add_heading("Certifications", level=2)
    doc.add_paragraph("Python for Everybody — Coursera")
    doc.add_paragraph("Introduction to Generative AI — Google Cloud")

    doc.add_heading("Achievements", level=2)
    doc.add_paragraph("Built and showcased 3 portfolio projects focused on AI and web development.")
    doc.add_paragraph("Participated in college hackathons and technical project exhibitions.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
